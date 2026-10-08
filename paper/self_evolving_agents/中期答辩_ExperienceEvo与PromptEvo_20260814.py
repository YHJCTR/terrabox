"""Generate the editable midterm-defense deck from the current manuscript package.

Run with: PYTHONPATH=/tmp/pptxgen python3 中期答辩_ExperienceEvo与PromptEvo_20260814.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("中期答辩_ExperienceEvo与PromptEvo_20260814.pptx")

# Academic palette: slate for shared framing, blue for ExperienceEvo,
# green for PromptEvo, and orange only for implementation-boundary warnings.
INK = "16212B"
MUTED = "5D6975"
LINE = "D7DEE5"
PANEL = "F5F7F9"
BLUE = "1769AA"
BLUE_PALE = "E8F2FA"
TEAL = "007F73"
TEAL_PALE = "E7F5F1"
ORANGE = "D97706"
ORANGE_PALE = "FFF3DE"
RED = "B42318"
WHITE = "FFFFFF"


def rgb(value):
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP,
             margin=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(0.8)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def box_text(slide, x, y, w, h, title, body="", fill=WHITE, accent=BLUE, title_size=16,
             body_size=11.5, border=LINE):
    add_box(slide, x, y, w, h, fill, border, radius=True)
    add_box(slide, x, y, 0.07, h, accent, accent)
    add_text(slide, x + 0.18, y + 0.13, w - 0.28, 0.3, title, title_size, INK, True)
    if body:
        add_text(slide, x + 0.18, y + 0.53, w - 0.30, h - 0.63, body, body_size, MUTED)


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=1.6):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    c.line.end_arrowhead = True
    return c


def add_header(slide, section, title, page, accent=BLUE):
    add_text(slide, 0.55, 0.24, 3.0, 0.25, section.upper(), 8.5, accent, True)
    add_text(slide, 0.55, 0.50, 11.6, 0.45, title, 27, INK, True)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.10), Inches(12.20), Inches(0.018))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(LINE); line.line.fill.background()
    add_text(slide, 12.35, 7.10, 0.35, 0.2, f"{page:02d}", 9, MUTED, False, align=PP_ALIGN.RIGHT)


def add_footer(slide, label="Terrabox | Self-Evolving Tool Agents"):
    add_text(slide, 0.55, 7.10, 5.5, 0.18, label, 8.5, MUTED)


def bullets(slide, x, y, w, lines, size=15, color=INK, gap=0.48):
    for i, line in enumerate(lines):
        add_text(slide, x, y + i * gap, 0.24, 0.28, "•", size + 2, BLUE, True)
        add_text(slide, x + 0.23, y + i * gap, w - 0.23, 0.34, line, size, color)


def base_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = rgb(WHITE)
    return slide


def title_slide(prs):
    s = base_slide(prs)
    add_box(s, 0, 0, 13.333, 0.14, BLUE, BLUE)
    add_text(s, 0.70, 0.86, 1.9, 0.28, "中期答辩", 13, TEAL, True)
    add_text(s, 0.70, 1.28, 11.8, 1.15, "面向工具增强 Agent 的\n可审计自进化", 34, INK, True)
    add_text(s, 0.73, 2.65, 10.8, 0.42, "ExperienceEvo：产品状态转移经验    |    PromptEvo：静态行为协议演化", 17, MUTED)
    add_text(s, 0.73, 3.22, 5.5, 0.32, "研究对象：多工具地理空间 Agent", 14, INK, True)
    add_text(s, 0.73, 3.62, 6.5, 0.32, "当前阶段：方法与实验协议已形成，正式证据正在冻结", 13, MUTED)
    # Two method cards.
    box_text(s, 0.75, 4.55, 5.65, 1.45, "ExperienceEvo", "从 rollout 提炼可复用的产品状态转移；在线以当前状态检索并给出软指导。", BLUE_PALE, BLUE, 20, 13)
    box_text(s, 6.93, 4.55, 5.65, 1.45, "PromptEvo", "从同任务对比轨迹提出小型静态协议补丁；经编译审计和开发集接受门筛选。", TEAL_PALE, TEAL, 20, 13)
    add_text(s, 0.75, 6.62, 9.0, 0.28, "答辩重点：方法机制、可证伪边界、实验闭环与下一步计划", 13, MUTED)
    add_footer(s)


def agenda_slide(prs):
    s = base_slide(prs); add_header(s, "Overview", "答辩路线", 2, TEAL)
    items = [
        ("01", "共同问题", "工具 Agent 的失败不是单一推理错误，而是状态、工具、参数、终止与评测的联动问题。", INK),
        ("02", "ExperienceEvo", "以产品状态转移为最小经验单元，研究可审计的 runtime reuse。", BLUE),
        ("03", "PromptEvo", "以 typed behavior protocol 为更新单元，研究可控的静态 prompt 自进化。", TEAL),
        ("04", "验证闭环", "冻结配置、分解对照、任务级审计、负结果保留与投稿门槛。", ORANGE),
    ]
    for i, (num, title, body, color) in enumerate(items):
        y = 1.48 + i * 1.22
        add_text(s, 0.8, y, 0.66, 0.42, num, 21, color, True)
        add_box(s, 1.55, y - 0.04, 10.65, 0.82, PANEL, LINE, radius=True)
        add_text(s, 1.82, y + 0.10, 2.2, 0.25, title, 16, INK, True)
        add_text(s, 4.1, y + 0.10, 7.6, 0.34, body, 12.5, MUTED)
    add_footer(s)


def problem_slide(prs):
    s = base_slide(prs); add_header(s, "Shared Motivation", "问题：多工具 Agent 缺少可审计的经验闭环", 3, TEAL)
    box_text(s, 0.72, 1.48, 3.75, 1.48, "状态断裂", "历史轨迹中的产物、当前运行的 artifact 与下游参数之间经常没有显式契约。", "F8FAFC", BLUE, 17, 12)
    box_text(s, 4.79, 1.48, 3.75, 1.48, "行为漂移", "全量重写 prompt 难以定位哪条规则导致回归，也难以阻止任务特异信息进入提示词。", "F8FAFC", TEAL, 17, 12)
    box_text(s, 8.86, 1.48, 3.75, 1.48, "证据混淆", "工具分数、最终答案、生成产物和服务异常常被混成单一指标。", "F8FAFC", ORANGE, 17, 12)
    add_text(s, 0.82, 3.62, 11.5, 0.34, "核心立场：自进化不应等同于“保存成功轨迹”或“改写一段 prompt”。", 20, INK, True, align=PP_ALIGN.CENTER)
    add_box(s, 1.15, 4.42, 10.95, 1.25, WHITE, LINE, radius=True)
    add_text(s, 1.48, 4.69, 10.3, 0.24, "需要同时回答三个问题", 14, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.56, 5.12, 3.0, 0.26, "何种历史信息可复用？", 15, INK, True, align=PP_ALIGN.CENTER)
    add_text(s, 5.20, 5.12, 3.0, 0.26, "如何在当前任务中安全使用？", 15, INK, True, align=PP_ALIGN.CENTER)
    add_text(s, 8.86, 5.12, 2.7, 0.26, "如何证明收益来源？", 15, INK, True, align=PP_ALIGN.CENTER)
    arrow(s, 4.65, 5.25, 5.00, 5.25, LINE); arrow(s, 8.34, 5.25, 8.70, 5.25, LINE)
    add_text(s, 0.92, 6.36, 11.5, 0.38, "两项工作分别处理 runtime 经验复用与 static behavior policy 更新，并在实验层统一采用冻结证据链。", 13, MUTED, align=PP_ALIGN.CENTER)
    add_footer(s)


def landscape_slide(prs):
    s = base_slide(prs); add_header(s, "Related Work", "研究定位：从经验存储和 prompt 优化走向可审计更新", 4, TEAL)
    cols = [
        (0.72, "经验/技能学习", "ExpeL, MemRL, AutoSkill, SkillRL", "通常强调记忆、技能或强化学习；本工作额外显式表示产物状态、绑定和 provenance。", BLUE_PALE, BLUE),
        (4.55, "工具 Agent", "ReAct, ToolLLM, OpenEarthAgent", "通常关注规划、调用与工具覆盖；本工作聚焦当前产品是否足以支持下一步。", "F5F7F9", "657786"),
        (8.38, "Prompt 优化", "APE, OPRO, ProTeGi, GEPA, SPRIG", "通常搜索完整 prompt 或候选文本；本工作将更新限制为 typed protocol patches，并单独定义接受门。", TEAL_PALE, TEAL),
    ]
    for x, title, refs, body, fill, accent in cols:
        box_text(s, x, 1.55, 3.42, 3.55, title, refs + "\n\n" + body, fill, accent, 17, 12)
    add_box(s, 1.32, 5.74, 10.7, 0.75, ORANGE_PALE, ORANGE, radius=True)
    add_text(s, 1.58, 5.98, 10.1, 0.25, "差异不在于“又增加一个 memory/prompt”，而在于把 provenance、接受条件、失败分母和反事实控制写进方法对象。", 13, INK, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def divider(prs, n, title, subtitle, color):
    s = base_slide(prs)
    add_box(s, 0, 0, 13.333, 7.5, color, color)
    add_text(s, 0.78, 1.30, 1.0, 0.36, f"PART {n}", 16, WHITE, True)
    add_text(s, 0.78, 1.95, 11.2, 0.70, title, 33, WHITE, True)
    add_text(s, 0.80, 2.90, 10.5, 0.35, subtitle, 16, "F2F7FA")
    add_box(s, 0.80, 4.25, 11.6, 0.02, "BFE1F3", "BFE1F3")
    add_text(s, 0.80, 4.60, 8.4, 0.50, "目标：把有效经验表达为可验证的状态转移，而不是把历史动作序列直接复制到新任务。", 18, WHITE)
    add_footer(s, "Terrabox | Self-Evolving Tool Agents")


def experience_overview(prs):
    s = base_slide(prs); add_header(s, "ExperienceEvo", "ExperienceEvo：从 rollout 到当前产品状态的软指导", 6, BLUE)
    # Offline lane
    add_text(s, 0.73, 1.43, 2.8, 0.25, "OFFLINE: train rollouts", 10, BLUE, True)
    box_text(s, 0.73, 1.78, 2.25, 1.15, "Event records", "tool call + input/output product + local evidence", BLUE_PALE, BLUE, 14, 10.5)
    box_text(s, 3.38, 1.78, 2.25, 1.15, "Family builder", "anonymize, group, update Q / N / R", BLUE_PALE, BLUE, 15, 10.5)
    box_text(s, 6.03, 1.78, 2.25, 1.15, "Read-only store", "product contract + alternative tool policies", BLUE_PALE, BLUE, 15, 10.5)
    arrow(s, 2.98, 2.35, 3.36, 2.35, BLUE); arrow(s, 5.63, 2.35, 6.01, 2.35, BLUE)
    # Online lane
    add_text(s, 0.73, 3.55, 2.8, 0.25, "ONLINE: current rollout", 10, TEAL, True)
    box_text(s, 0.73, 3.90, 2.20, 1.12, "Live state", "current images, data, artifacts, available tools", TEAL_PALE, TEAL, 15, 10.5)
    box_text(s, 3.38, 3.90, 2.20, 1.12, "Retrieve", "state match + fixed heuristic + policy rank", TEAL_PALE, TEAL, 15, 10.5)
    box_text(s, 6.03, 3.90, 2.20, 1.12, "Soft guidance", "current-run binding, output checks, recovery", TEAL_PALE, TEAL, 15, 10.5)
    box_text(s, 8.68, 3.90, 2.20, 1.12, "Base agent", "chooses action; guidance is never a forced script", TEAL_PALE, TEAL, 15, 10.5)
    arrow(s, 2.93, 4.46, 3.36, 4.46, TEAL); arrow(s, 5.58, 4.46, 6.01, 4.46, TEAL); arrow(s, 8.23, 4.46, 8.66, 4.46, TEAL)
    arrow(s, 7.10, 2.95, 5.00, 3.85, BLUE)
    # warning
    box_text(s, 9.38, 1.78, 3.15, 1.15, "Implementation boundary", "v3 deterministic fallback may return before learned-family ranking. It is not rollout-derived experience.", ORANGE_PALE, ORANGE, 14, 10.2, ORANGE)
    arrow(s, 10.75, 2.96, 4.80, 3.84, ORANGE, 1.2)
    add_text(s, 0.75, 5.75, 11.8, 0.54, "研究假设：相对于相同 base agent，状态兼容的经验能够减少无效调用，并提高可追踪的产品链与请求产物完成度。", 16, INK, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def experience_unit(prs):
    s = base_slide(prs); add_header(s, "ExperienceEvo", "经验单元：产品契约，而不是历史动作脚本", 7, BLUE)
    add_text(s, 0.78, 1.42, 11.6, 0.35, "一个 transition family 以“当前已有何种可用产品 -> 下一步需要产生何种产品”为核心。", 18, INK, True)
    stages = [
        (0.82, "Input state", "已有 image / data-file / layer / measurement", BLUE_PALE, BLUE),
        (3.33, "Product contract", "preconditions | output checks | current-run binding | recovery", "F5F7F9", "657786"),
        (6.92, "Tool policies", "同一目标可保留多个 tool route，而不是固定一条历史序列", TEAL_PALE, TEAL),
        (10.43, "Target state", "new valid layer / count / map / calculation / annotation", BLUE_PALE, BLUE),
    ]
    widths = [2.05, 3.10, 3.05, 2.05]
    for i, (x, title, body, fill, accent) in enumerate(stages):
        box_text(s, x, 2.05, widths[i], 1.62, title, body, fill, accent, 16, 11)
        if i < 3:
            arrow(s, x + widths[i], 2.85, stages[i + 1][0] - 0.05, 2.85, accent)
    add_box(s, 1.10, 4.50, 11.10, 1.22, WHITE, LINE, radius=True)
    add_text(s, 1.38, 4.75, 10.55, 0.27, "局部证据", 15, BLUE, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.40, 5.15, 10.45, 0.27, "input binding valid  |  output valid  |  target completed  |  downstream consumed  |  attributable risk", 13, MUTED, align=PP_ALIGN.CENTER)
    add_text(s, 1.10, 6.22, 11.05, 0.32, "边界：$Q/R$ 是离线平滑的局部经验统计，不是最终任务成功的因果归因或校准概率。", 13, ORANGE, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def experience_runtime(prs):
    s = base_slide(prs); add_header(s, "ExperienceEvo", "在线机制：状态匹配、当前引用绑定与软验证", 8, BLUE)
    steps = [
        ("1", "State", "从当前 image、data、artifact tracker 推得 live product state。", BLUE),
        ("2", "Retrieve", "筛选前置状态兼容且工具可用的 transition families。", BLUE),
        ("3", "Guide", "注入短规则：使用当前返回值、检查输出、失败后恢复。", TEAL),
        ("4", "Act", "base agent 自主选 tool；hard guard 仅处理窄范围 schema/path。", TEAL),
        ("5", "Update", "根据新 observation 更新状态，按需提供下一步 hint。", BLUE),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        x = 0.65 + i * 2.48
        add_box(s, x, 1.83, 2.08, 2.25, WHITE, LINE, radius=True)
        add_text(s, x + 0.18, 2.05, 0.40, 0.35, num, 20, color, True)
        add_text(s, x + 0.63, 2.07, 1.25, 0.28, title, 16, INK, True)
        add_text(s, x + 0.18, 2.66, 1.70, 1.05, body, 11.5, MUTED)
        if i < 4:
            arrow(s, x + 2.10, 2.96, x + 2.44, 2.96, color)
    add_box(s, 0.95, 4.78, 11.55, 1.18, ORANGE_PALE, ORANGE, radius=True)
    add_text(s, 1.22, 5.05, 10.95, 0.30, "必须单独审计 fallback", 15, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.22, 5.45, 10.95, 0.26, "当前实现中，nonempty v3 fallback 会先返回，可能遮蔽 ordinary learned-store ranking；其收益不能归给 train-rollout experience。", 12.3, INK, align=PP_ALIGN.CENTER)
    add_footer(s)


def experience_eval(prs):
    s = base_slide(prs); add_header(s, "ExperienceEvo", "验证设计：先分解工程保护，再谈经验迁移", 9, BLUE)
    cols = [
        ("Base", "无 store\n无 fallback\n无 guard", "F8FAFC", "657786"),
        ("+ Generic guard", "仅 image binding\n+ calculator syntax", ORANGE_PALE, ORANGE),
        ("No-store soft-only", "无经验库\n保留软状态/验证语言", "F8FAFC", "657786"),
        ("Fallback-only", "无 learned store\n仅 deterministic fallback", ORANGE_PALE, ORANGE),
        ("Learned-store-only", "read-only store\nfallback disabled", BLUE_PALE, BLUE),
        ("Full v4", "store + fallback\n+ guard + verifier", TEAL_PALE, TEAL),
    ]
    for i, (title, body, fill, accent) in enumerate(cols):
        x = 0.45 + i * 2.11
        box_text(s, x, 1.62, 1.82, 1.45, title, body, fill, accent, 12.2, 10.1)
        if i < 5: arrow(s, x + 1.84, 2.34, x + 2.07, 2.34, LINE, 1.0)
    add_text(s, 0.73, 3.68, 11.8, 0.30, "主要指标必须拆开报告", 16, INK, True, align=PP_ALIGN.CENTER)
    metrics = [("执行", "Success / tool agreement / invalid & repeated calls", BLUE), ("答案", "ordinary answer quality with independent audit", TEAL), ("产物", "requested map/file/annotation contract completion", ORANGE), ("成本", "tools, LLM calls, tokens, latency, failures", "657786")]
    for i, (title, body, color) in enumerate(metrics):
        x = 0.78 + i * 3.00
        add_box(s, x, 4.25, 2.62, 1.14, WHITE, LINE, radius=True)
        add_text(s, x + 0.12, 4.50, 2.38, 0.22, title, 14, color, True, align=PP_ALIGN.CENTER)
        add_text(s, x + 0.15, 4.85, 2.32, 0.32, body, 9.7, MUTED, align=PP_ALIGN.CENTER)
    add_text(s, 0.82, 6.22, 11.6, 0.33, "正式 mechanism log：family ID、provenance、fallback suppression、recommended/selected tool、binding/product checks。", 13, MUTED, align=PP_ALIGN.CENTER)
    add_footer(s)


def prompt_divider(prs):
    s = base_slide(prs)
    add_box(s, 0, 0, 13.333, 7.5, TEAL, TEAL)
    add_text(s, 0.78, 1.30, 1.0, 0.36, "PART 02", 16, WHITE, True)
    add_text(s, 0.78, 1.95, 11.2, 0.70, "PromptEvo：静态行为协议的对比式自进化", 30, WHITE, True)
    add_text(s, 0.80, 2.90, 10.5, 0.35, "Contrastive self-evolution of static behavior protocols", 16, "F2F7FA")
    add_box(s, 0.80, 4.25, 11.6, 0.02, "BDEBE5", "BDEBE5")
    add_text(s, 0.80, 4.60, 8.9, 0.50, "目标：把 prompt 更新限制为可审计的小型行为规则，并把“候选生成”与“真实 rollout 接受”明确分开。", 18, WHITE)
    add_footer(s)


def prompt_loop(prs):
    s = base_slide(prs); add_header(s, "PromptEvo", "PromptEvo：从 paired traces 到经接受的静态协议", 11, TEAL)
    nodes = [
        (0.62, 2.02, 2.02, "A / B prompts", "same-task rollout traces\n+ literal prompt diff", TEAL_PALE, TEAL),
        (3.12, 2.02, 2.02, "Attribution", "separate prompt-fixable cases\nfrom tool/service noise", "F5F7F9", "657786"),
        (5.62, 2.02, 2.02, "Typed patches", "trigger + rule + kind\n+ priority + risk + evidence", TEAL_PALE, TEAL),
        (8.12, 2.02, 2.02, "Compiler audit", "reject task-specific text, conflicts; preserve anchors", "F5F7F9", "657786"),
        (10.62, 2.02, 2.02, "Frozen dev", "primary + protected metrics\naccept or reject", TEAL_PALE, TEAL),
    ]
    for i, (x, y, w, title, body, fill, accent) in enumerate(nodes):
        box_text(s, x, y, w, 1.54, title, body, fill, accent, 14, 10.2)
        if i < 4: arrow(s, x + w, y + 0.77, nodes[i + 1][0] - 0.05, y + 0.77, TEAL)
    add_box(s, 4.20, 4.66, 4.92, 1.05, WHITE, LINE, radius=True)
    add_text(s, 4.45, 4.92, 4.42, 0.28, "Locked test: one frozen fingerprint", 15, INK, True, align=PP_ALIGN.CENTER)
    add_text(s, 4.45, 5.30, 4.42, 0.20, "never used to revise thresholds or patches", 10.8, MUTED, align=PP_ALIGN.CENTER)
    arrow(s, 11.62, 3.62, 8.95, 4.64, TEAL)
    add_text(s, 0.84, 6.29, 11.6, 0.29, "关键区分：静态 analysis 可排序候选；只有在冻结 dev 上通过预声明指标的 rollout，才可称为 accepted。", 13, ORANGE, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def prompt_patch(prs):
    s = base_slide(prs); add_header(s, "PromptEvo", "更新对象：typed behavior protocol patches", 12, TEAL)
    kinds = [
        ("Tool selection", "在工具有歧义时先核对当前目标和可用输入。", BLUE),
        ("Argument validation", "使用 artifact-like 参数前，确认其来自当前运行且满足 schema。", TEAL),
        ("Error recovery", "遇到可归因的参数/输出错误时，先恢复缺失产品再重试。", ORANGE),
        ("Termination & repetition", "仅在请求证据完整时结束；避免重复同一失败绑定。", "657786"),
    ]
    for i, (title, example, accent) in enumerate(kinds):
        y = 1.52 + i * 1.08
        add_box(s, 0.78, y, 11.80, 0.78, WHITE, LINE, radius=True)
        add_text(s, 1.05, y + 0.22, 2.30, 0.24, title, 14, accent, True)
        add_text(s, 3.56, y + 0.19, 8.55, 0.31, example, 12.5, MUTED)
    add_box(s, 1.00, 6.05, 11.3, 0.56, ORANGE_PALE, ORANGE, radius=True)
    add_text(s, 1.22, 6.22, 10.85, 0.22, "Compiler guarantees syntactic/policy-surface constraints, not semantic safety, causal utility, or cross-environment generalization。", 11.5, INK, align=PP_ALIGN.CENTER)
    add_footer(s)


def prompt_boundary(prs):
    s = base_slide(prs); add_header(s, "PromptEvo", "实现边界：当前代码与论文协议必须严格区分", 13, TEAL)
    headers = [(0.78, "当前已实现", TEAL_PALE, TEAL), (4.52, "当前尚未自动保证", ORANGE_PALE, ORANGE), (8.26, "正式提交前的证据", BLUE_PALE, BLUE)]
    for x, title, fill, accent in headers:
        add_box(s, x, 1.48, 3.22, 0.55, fill, accent, radius=True)
        add_text(s, x + 0.15, 1.64, 2.92, 0.20, title, 13.5, accent, True, align=PP_ALIGN.CENTER)
    left = ["typed patch kinds", "task/path pattern checks", "conflict / duplicate checks", "high-risk withholding", "current success-drop acceptance hook"]
    mid = ["evidence nonempty", "evidence resolves to trace/attribution", "all protected metrics", "candidate ledger linkage", "CLI persistence equals acceptance"]
    right = ["all candidates + rejections", "frozen dev split hash", "thresholds and paired deltas", "selected prompt fingerprint", "one locked-test result"]
    for col, lines, accent in [(0.94, left, TEAL), (4.68, mid, ORANGE), (8.42, right, BLUE)]:
        for i, line in enumerate(lines):
            add_text(s, col, 2.31 + i * 0.59, 0.18, 0.22, "•", 14, accent, True)
            add_text(s, col + 0.18, 2.31 + i * 0.59, 2.77, 0.28, line, 12.2, INK)
    add_text(s, 0.92, 6.18, 11.4, 0.36, "因此：被写入文件的 prompt 只能叫 persisted candidate；没有真实 dev 决策与证据链，不能叫 accepted prompt。", 13, ORANGE, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def prompt_eval(prs):
    s = base_slide(prs); add_header(s, "PromptEvo", "实验设计：把搜索过程本身变成可复核对象", 14, TEAL)
    stages = [
        ("Construction", "freeze paired tasks, base/current fingerprints, literal diff, trace batches", "F5F7F9", "657786"),
        ("Candidate audit", "compiler report + patch-level declared evidence + resolution status", TEAL_PALE, TEAL),
        ("Dev acceptance", "predeclare primary/protected metrics, thresholds, failure policy, ledger", BLUE_PALE, BLUE),
        ("Locked test", "run exactly one frozen accepted fingerprint; preserve negative outcomes", ORANGE_PALE, ORANGE),
    ]
    for i, (title, body, fill, accent) in enumerate(stages):
        x = 0.78 + i * 3.12
        box_text(s, x, 1.72, 2.70, 2.05, title, body, fill, accent, 15, 11)
        if i < 3: arrow(s, x + 2.72, 2.75, x + 3.05, 2.75, accent)
    add_box(s, 1.05, 4.58, 11.15, 1.28, WHITE, LINE, radius=True)
    add_text(s, 1.32, 4.86, 10.6, 0.26, "no-update control 必须与每一个 prompt 条件共享 actor、tool/evaluator manifest、retry/timeout/cache 和 worker policy。", 13.5, INK, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.32, 5.30, 10.6, 0.24, "报告 utility、安全/contract、成本、基础设施失败和候选数；不以加权总分掩盖显著回归。", 12.2, MUTED, align=PP_ALIGN.CENTER)
    add_footer(s)


def evidence_slide(prs):
    s = base_slide(prs); add_header(s, "Evidence Plan", "统一证据链：主张必须能回到任务级记录", 15, ORANGE)
    layers = [
        ("Raw results", "每个计划 task 均保留：completed / failed / empty final / evaluator error。", "F8FAFC", "657786"),
        ("Immutable manifest", "split、模型、prompt、工具 schema、cache/retry、service/resource policy。", BLUE_PALE, BLUE),
        ("Mechanism record", "ExperienceEvo provenance trace；PromptEvo candidate ledger / compiler report。", TEAL_PALE, TEAL),
        ("Paired analysis", "共同 task IDs、分母、bootstrap CI、失败核算、分层分析。", ORANGE_PALE, ORANGE),
        ("Paper claim", "主表、图、摘要与结论只使用已冻结证据支持的范围。", "F8FAFC", "657786"),
    ]
    for i, (title, body, fill, accent) in enumerate(layers):
        y = 1.38 + i * 0.93
        add_box(s, 1.45, y, 10.42, 0.65, fill, accent, radius=True)
        add_text(s, 1.72, y + 0.17, 2.1, 0.20, title, 13.4, accent, True)
        add_text(s, 3.98, y + 0.16, 7.45, 0.24, body, 11.5, MUTED)
        if i < 4: arrow(s, 6.66, y + 0.67, 6.66, y + 0.89, LINE, 1.1)
    add_text(s, 1.20, 6.37, 11.0, 0.25, "历史诊断用于设计与风险识别，不替代 final evidence；任何主表数字必须可追溯。", 13, ORANGE, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def current_status(prs):
    s = base_slide(prs); add_header(s, "Current Status", "当前进度：方法稿完成，正式论文证据尚未完成", 16, ORANGE)
    box_text(s, 0.78, 1.55, 3.55, 2.38, "已完成", "两篇独立主稿与补充材料\n方法图、相关工作、实现审计\n实验矩阵、证据包模板、投稿检查单\n历史诊断与负结果保留", TEAL_PALE, TEAL, 18, 12)
    box_text(s, 4.90, 1.55, 3.55, 2.38, "进行中", "冻结 train/dev/test 与 manifests\n补齐 ExperienceEvo provenance trace\n闭合 PromptEvo evidence/acceptance ledger\n完成主比较和必要消融", BLUE_PALE, BLUE, 18, 12)
    box_text(s, 9.02, 1.55, 3.55, 2.38, "提交前", "official/adapted baseline fidelity\npaired uncertainty + independent audit\n填主表并复核所有 claim\nOverleaf PDF、作者、引用、许可", ORANGE_PALE, ORANGE, 18, 12)
    add_box(s, 1.05, 4.77, 11.18, 0.96, WHITE, LINE, radius=True)
    add_text(s, 1.35, 5.02, 10.58, 0.28, "当前不应宣称：最终性能提升、SOTA、经验迁移的因果收益、prompt 的安全非退化或跨环境泛化。", 13, RED, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.35, 5.42, 10.58, 0.18, "当前可陈述：两套可审计方法对象、明确实现边界，以及可证伪的实验协议。", 12, MUTED, align=PP_ALIGN.CENTER)
    add_footer(s)


def next_steps(prs):
    s = base_slide(prs); add_header(s, "Plan", "下一阶段：从方法稿走向可提交的证据包", 17, ORANGE)
    timeline = [
        ("1", "冻结研究对象", "split、模型、工具、prompt 与运行策略。", BLUE),
        ("2", "补齐可审计日志", "ExperienceEvo provenance；PromptEvo resolution 与 ledger。", TEAL),
        ("3", "运行关键对照", "causal matrix；no-update；frozen dev / locked test。", ORANGE),
        ("4", "统计与审计", "paired CI、failure accounting、answer/artifact audit。", BLUE),
        ("5", "论文冻结", "填表、claim audit、PDF、许可与 arXiv。", TEAL),
    ]
    for i, (num, title, body, color) in enumerate(timeline):
        x = 0.52 + i * 2.54
        add_text(s, x + 0.79, 1.58, 0.55, 0.44, num, 25, color, True, align=PP_ALIGN.CENTER)
        add_box(s, x, 2.15, 2.13, 2.22, WHITE, LINE, radius=True)
        add_text(s, x + 0.12, 2.48, 1.89, 0.30, title, 13.5, color, True, align=PP_ALIGN.CENTER)
        add_text(s, x + 0.14, 3.08, 1.84, 0.94, body, 9.4, MUTED, align=PP_ALIGN.CENTER)
        if i < 4: arrow(s, x + 2.16, 3.25, x + 2.49, 3.25, LINE, 1.1)
    add_text(s, 1.15, 5.63, 11.0, 0.45, "成功标准不是“跑出更高数字”，而是让任何核心结论都能回溯到同一冻结配置下的原始任务级证据。", 15, INK, True, align=PP_ALIGN.CENTER)
    add_footer(s)


def closing(prs):
    s = base_slide(prs)
    add_box(s, 0, 0, 13.333, 7.5, INK, INK)
    add_text(s, 0.78, 1.22, 1.8, 0.25, "TAKEAWAYS", 11, "8AC5D1", True)
    add_text(s, 0.78, 1.72, 11.0, 0.68, "将 Agent 自进化从“经验堆积”\n转化为可审计的更新闭环", 30, WHITE, True)
    takes = [
        ("ExperienceEvo", "用产品状态转移表达 runtime experience reuse，并把 fallback 与 learned store 严格分开。", BLUE),
        ("PromptEvo", "用 typed protocol patches 约束静态行为更新，并把生成、编译、接受和测试分开。", TEAL),
        ("Evidence first", "冻结配置、反事实控制和任务级 provenance 决定论文能否成立。", ORANGE),
    ]
    for i, (title, body, color) in enumerate(takes):
        y = 4.00 + i * 0.62
        add_text(s, 0.94, y, 2.1, 0.25, title, 14, color, True)
        add_text(s, 3.10, y, 8.8, 0.28, body, 12.5, "E5EDF1")
    add_text(s, 0.78, 6.78, 11.0, 0.25, "Thank you", 13, "8AC5D1")


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "面向工具增强 Agent 的可审计自进化"
    prs.core_properties.subject = "ExperienceEvo 与 PromptEvo 中期答辩"
    prs.core_properties.author = "Terrabox"
    title_slide(prs)
    agenda_slide(prs)
    problem_slide(prs)
    landscape_slide(prs)
    divider(prs, "01", "ExperienceEvo：产品状态转移经验", "Product-State Transition Self-Evolution for Tool-Augmented Geospatial Agents", BLUE)
    experience_overview(prs)
    experience_unit(prs)
    experience_runtime(prs)
    experience_eval(prs)
    prompt_divider(prs)
    prompt_loop(prs)
    prompt_patch(prs)
    prompt_boundary(prs)
    prompt_eval(prs)
    evidence_slide(prs)
    current_status(prs)
    next_steps(prs)
    closing(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
